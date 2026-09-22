from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from styles import (
    F_HEADER_LIGHT, F_BODY, FILL_HEADER, sheet_header, set_sheet_defaults, F_SECTION,
    ALIGN_LEFT, ALIGN_CENTER, FMT_CURRENCY0, FMT_DATE, FMT_INT, col_idx, DUSTY_ROSE,
    FILL_CREAM, B_CARD, F_KPI_LABEL,
)

LISTS = {
    "RSVPStatus": ["Pending", "Attending", "Declined", "Tentative"],
    "PaymentStatus": ["Not Started", "Deposit Paid", "Partially Paid", "Paid in Full", "Overdue"],
    "TaskStatus": ["Not Started", "In Progress", "Completed", "Overdue"],
    "Priority": ["Low", "Medium", "High"],
    "ContractStatus": ["Not Sent", "Sent", "Signed", "Cancelled"],
    "BudgetCategory": ["Venue", "Catering", "Photography", "Videography", "Attire", "Beauty",
                        "Flowers", "Decor", "Entertainment", "Music", "Transportation",
                        "Invitations", "Stationery", "Cake", "Rings", "Rentals",
                        "Accommodation", "Gifts", "Honeymoon", "Miscellaneous"],
    "MealPreference": ["Standard", "Vegetarian", "Vegan", "Gluten-Free", "Kids Meal", "Other"],
    "InvitationStatus": ["Not Sent", "Save-the-Date Sent", "Invitation Sent"],
    "Side": ["Partner 1", "Partner 2", "Both"],
    "Owner": ["Partner 1", "Partner 2", "Wedding Planner", "Both", "Parents"],
    "PaymentMethod": ["Bank Transfer", "Credit Card", "Cash", "Cheque", "Digital Wallet"],
    "YesNo": ["Yes", "No"],
    "GiftType": ["Physical Gift", "Cash Gift", "Gift Card", "Experience", "Registry Item"],
    "ChecklistPeriod": ["12+ Months Before", "9-12 Months Before", "6-9 Months Before",
                         "3-6 Months Before", "1-3 Months Before", "Final Month",
                         "Wedding Week", "Wedding Day", "Post-Wedding"],
    "FoodCategory": ["Appetizers", "Main Course", "Dessert", "Cake", "Drinks", "Late Night Food"],
    "HoneymoonCategory": ["Flights", "Hotel", "Transportation", "Food", "Activities",
                           "Shopping", "Travel Insurance", "Miscellaneous"],
    "Rating": ["1", "2", "3", "4", "5"],
    "PlusOne": ["Yes", "No"],
    "Children": ["Yes", "No"],
}


def build_lists_sheet(wb):
    ws = wb.create_sheet("Lists")
    set_sheet_defaults(ws)
    ws.sheet_state = "hidden"
    ranges = {}
    for i, (key, values) in enumerate(LISTS.items()):
        c = i + 1
        letter = get_column_letter(c)
        ws.cell(row=1, column=c, value=key)
        for r, v in enumerate(values, start=2):
            ws.cell(row=r, column=c, value=v)
        ranges[key] = f"Lists!${letter}$2:${letter}${len(values) + 1}"
        ws.column_dimensions[letter].width = 20
    return ranges


def build_settings_sheet(wb, demo=True):
    ws = wb.create_sheet("Settings")
    set_sheet_defaults(ws)
    row = sheet_header(ws, "SETTINGS", "Centralize your wedding details, budget and preferences once — every other sheet updates automatically.", "F")
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 4
    ws.column_dimensions["E"].width = 26
    ws.column_dimensions["F"].width = 30

    fields = [
        ("Partner 1 Name", "Sophia Bennett" if demo else ""),
        ("Partner 2 Name", "Daniel Whitfield" if demo else ""),
        ("Wedding Date", "2027-06-15" if demo else ""),
        ("Currency Symbol", "$"),
        ("Wedding Location", "Napa Valley, California" if demo else ""),
        ("Guest Target", 150 if demo else ""),
        ("Total Wedding Budget", 70000 if demo else ""),
        ("Number of Tables", 16),
        ("Seats per Table", 10),
        ("Honeymoon Budget", 8000 if demo else ""),
        ("RSVP Reply-By Date", "2027-04-15" if demo else ""),
    ]
    r = row + 1
    cell_map = {}
    for label, value in fields:
        ws.cell(row=r, column=2, value=label).font = F_KPI_LABEL
        vcell = ws.cell(row=r, column=3, value=value)
        vcell.font = F_BODY
        vcell.fill = FILL_CREAM
        vcell.alignment = ALIGN_LEFT
        if label == "Wedding Date" or "Date" in label:
            vcell.number_format = FMT_DATE
        if "Budget" in label:
            vcell.number_format = FMT_CURRENCY0
        cell_map[label] = f"Settings!$C${r}"
        r += 1

    note_row = r + 1
    ws.cell(row=note_row, column=2,
            value="Edit the values above — the Dashboard, Budget, Guest List and every other sheet reference these settings automatically.").font = Font(
        name="Calibri", size=9.5, italic=True, color="715B52")
    ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=6)

    # Named ranges for convenience
    add_named_range(wb, "Partner1Name", cell_map["Partner 1 Name"])
    add_named_range(wb, "Partner2Name", cell_map["Partner 2 Name"])
    add_named_range(wb, "WeddingDate", cell_map["Wedding Date"])
    add_named_range(wb, "CurrencySymbol", cell_map["Currency Symbol"])
    add_named_range(wb, "TotalBudget", cell_map["Total Wedding Budget"])
    add_named_range(wb, "GuestTarget", cell_map["Guest Target"])
    add_named_range(wb, "TableCount", cell_map["Number of Tables"])
    add_named_range(wb, "SeatsPerTable", cell_map["Seats per Table"])
    add_named_range(wb, "HoneymoonBudgetSetting", cell_map["Honeymoon Budget"])
    add_named_range(wb, "RSVPDeadline", cell_map["RSVP Reply-By Date"])

    ws.page_setup.orientation = "portrait"
    return cell_map


def add_named_range(wb, name, ref):
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names[name] = DefinedName(name=name, attr_text=ref)
