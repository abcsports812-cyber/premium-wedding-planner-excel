"""Generic data-table sheet builder shared by Budget, Vendors, Guest List,
Payments, Checklist, Contacts, Gifts, Wedding Party, Venue, Food & Drinks,
Honeymoon Budget."""

from dataclasses import dataclass, field
from typing import Callable, Optional
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

from styles import (
    F_HEADER_LIGHT, F_BODY, FILL_HEADER, FILL_WHITE, ROW_B, fill, B_ALL_LIGHT,
    ALIGN_LEFT, ALIGN_LEFT_WRAP, ALIGN_CENTER, ALIGN_RIGHT, sheet_header,
    set_sheet_defaults, FMT_CURRENCY, FMT_DATE, FMT_INT, FMT_PCT, col_idx,
    LIGHT_SAGE, LIGHT_BLUSH, MUTED_GOLD, DEEP_ROSE, SAGE, DUSTY_ROSE, GOLD, CHARCOAL,
)


@dataclass
class Col:
    name: str
    width: float = 14
    number_format: Optional[str] = None
    validation_list: Optional[str] = None  # formula range string e.g. "Lists!$A$2:$A$5"
    formula: Optional[Callable[[int], str]] = None  # f(row) -> formula string (no leading '=')
    align: str = "left"
    wrap: bool = False


def build_table_sheet(wb, sheet_name, title, subtitle, columns, demo_rows, table_name,
                       n_formula_rows, demo, nav_links=None, tab_color=DUSTY_ROSE,
                       start_col=2, print_landscape=True, extra_after_header=None):
    """Creates a worksheet with banner + an Excel Table.
    columns: list[Col]. demo_rows: list[dict[name->value]] (only used when demo=True).
    n_formula_rows: total data rows to pre-provision with formulas (>= len(demo_rows)).
    Returns (ws, header_row, first_data_row, last_data_row, last_col_letter).
    """
    ws = wb.create_sheet(sheet_name)
    set_sheet_defaults(ws, tab_color)
    last_col = start_col + len(columns) - 1
    last_col_letter = get_column_letter(last_col)
    first_free_row = sheet_header(ws, title, subtitle, last_col_letter, nav_links)
    row = first_free_row
    if extra_after_header:
        row = extra_after_header(ws, row, last_col_letter) or row

    header_row = row
    for i, col in enumerate(columns):
        c = start_col + i
        cell = ws.cell(row=header_row, column=c, value=col.name)
        cell.font = F_HEADER_LIGHT
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        ws.column_dimensions[get_column_letter(c)].width = col.width

    n_rows = max(n_formula_rows, len(demo_rows), 1)
    first_data_row = header_row + 1
    last_data_row = first_data_row + n_rows - 1

    for r_i in range(n_rows):
        row_n = first_data_row + r_i
        demo_data = demo_rows[r_i] if demo and r_i < len(demo_rows) else {}
        for i, col in enumerate(columns):
            c = start_col + i
            cell = ws.cell(row=row_n, column=c)
            if col.formula:
                cell.value = "=" + col.formula(row_n)
            elif col.name in demo_data:
                cell.value = demo_data[col.name]
            if col.number_format:
                cell.number_format = col.number_format
            cell.font = F_BODY
            cell.border = B_ALL_LIGHT
            if col.align == "center":
                cell.alignment = ALIGN_CENTER_or_wrap(col)
            elif col.align == "right":
                cell.alignment = ALIGN_RIGHT
            else:
                cell.alignment = ALIGN_LEFT_WRAP if col.wrap else ALIGN_LEFT
        ws.row_dimensions[row_n].height = 20

    # Excel Table
    ref = f"{get_column_letter(start_col)}{header_row}:{last_col_letter}{last_data_row}"
    tab = Table(displayName=table_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showFirstColumn=False,
                                         showLastColumn=False, showRowStripes=True,
                                         showColumnStripes=False)
    ws.add_table(tab)

    # data validation dropdowns
    for i, col in enumerate(columns):
        if col.validation_list:
            c = start_col + i
            dv = DataValidation(type="list", formula1=col.validation_list, allow_blank=True,
                                 showDropDown=False)
            dv.error = "Please choose a value from the list."
            dv.errorTitle = "Invalid entry"
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(c)}{first_data_row}:{get_column_letter(c)}{last_data_row}")

    ws.freeze_panes = f"{get_column_letter(start_col)}{first_data_row}"
    ws.auto_filter.ref = ref

    # print setup
    ws.page_setup.orientation = "landscape" if print_landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_area = f"A1:{last_col_letter}{last_data_row}"

    return ws, header_row, first_data_row, last_data_row, last_col_letter


def ALIGN_CENTER_or_wrap(col):
    from openpyxl.styles import Alignment
    return Alignment(horizontal="center", vertical="center", wrap_text=col.wrap)


def status_conditional_formatting(ws, rng, mapping):
    """mapping: dict[text-> hex fill]; applies fill when cell text equals key (case sensitive)."""
    for text, color in mapping.items():
        rule = FormulaRule(formula=[f'EXACT(TRIM({rng.split(":")[0]}),"{text}")'],
                            fill=fill(color))
        ws.conditional_formatting.add(rng, rule)
