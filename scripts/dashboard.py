from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from styles import (
    set_sheet_defaults, F_TITLE, F_SUBTITLE, F_SECTION, F_BODY, F_KPI_LABEL, kpi_card,
    FMT_CURRENCY0, FMT_PCT1, FMT_INT, FMT_DATE, col_idx, DEEP_ROSE, SAGE, GOLD, BLUSH,
    LIGHT_BLUSH, LIGHT_SAGE, MUTED_GOLD, CHARCOAL, DUSTY_ROSE, IVORY, FILL_LIGHT_BLUSH,
    fill, F_NAV,
)
from lists_settings import LISTS

LIST_COL = {k: i + 1 for i, k in enumerate(LISTS.keys())}

PALETTE = [DUSTY_ROSE, SAGE, GOLD, "D9B8B4", "C7D2C0", MUTED_GOLD, "9D646B", "715B52"]


def style_chart(chart, title, height=8.2, width=15.5):
    chart.title = title
    chart.style = 2
    chart.height = height
    chart.width = width
    if hasattr(chart, "y_axis"):
        chart.y_axis.majorGridlines = None
    if hasattr(chart, "gapWidth"):
        chart.gapWidth = 60
    for i, series in enumerate(chart.series):
        color = PALETTE[i % len(PALETTE)]
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.noFill = True
    return chart


def colorize_doughnut(chart, n):
    series = chart.series[0]
    pts = []
    for i in range(n):
        dp = DataPoint(idx=i)
        dp.graphicalProperties.solidFill = PALETTE[i % len(PALETTE)]
        pts.append(dp)
    series.data_points = pts


def build_dashboard(wb, refs, demo):
    """refs: dict with first/last rows for Budget, Vendors, Guest List, Payments,
    Master Checklist, Honeymoon Budget sheets."""
    ws = wb.create_sheet("Dashboard", 1)
    set_sheet_defaults(ws, GOLD)

    # ---- Wedding identity banner ----
    for r in range(1, 6):
        for c in range(1, 17):
            ws.cell(row=r, column=c).fill = FILL_LIGHT_BLUSH
    ws.row_dimensions[1].height = 6
    ws.row_dimensions[2].height = 16
    lbl = ws.cell(row=2, column=2, value="OUR WEDDING")
    lbl.font = Font(name="Calibri", size=11, bold=True, color=DEEP_ROSE)
    ws.row_dimensions[3].height = 34
    names = ws.cell(row=3, column=2, value='=IF(Partner1Name="","Partner One",Partner1Name)&"  &  "&IF(Partner2Name="","Partner Two",Partner2Name)')
    names.font = Font(name="Calibri", size=26, bold=True, color=CHARCOAL)
    ws.row_dimensions[4].height = 18
    datecell = ws.cell(row=4, column=2, value='=IF(WeddingDate="","Wedding date not yet set",TEXT(WeddingDate,"dddd, d mmmm yyyy"))')
    datecell.font = Font(name="Calibri", size=13, italic=True, color="715B52")
    for r in range(2, 5):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)

    ws.merge_cells("M2:P5")
    cd = ws.cell(row=2, column=13,
                 value='=IF(WeddingDate="","—",IF(WeddingDate>=TODAY(),WeddingDate-TODAY()&" DAYS TO GO",TEXT(WeddingDate,"d mmm yyyy")&" ❦ Married!"))')
    cd.font = Font(name="Calibri", size=17, bold=True, color=DEEP_ROSE)
    cd.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(2, 6):
        for c in range(13, 17):
            ws.cell(row=r, column=c).border = Border(left=Side(style="thin", color=GOLD))
    for c in range(1, 17):
        ws.cell(row=5, column=c).border = Border(bottom=Side(style="medium", color=GOLD))
    nav = ws.cell(row=1, column=2, value="‹ Start Here")
    nav.font = F_NAV
    nav.hyperlink = "#'START HERE'!A1"

    B, V, G, P, C, H = refs["Budget"], refs["Vendors"], refs["Guest List"], refs["Payments"], refs["Master Checklist"], refs["Honeymoon Budget"]

    row = 7
    sect = ws.cell(row=row, column=2, value="FINANCIAL SNAPSHOT")
    sect.font = F_SECTION
    row += 1
    cards_row = row
    kpi_card(ws, cards_row, 2, "Total Budget", "=TotalBudget", FMT_CURRENCY0, DEEP_ROSE)
    kpi_card(ws, cards_row, 5, "Actual Spending", f"=SUM(Budget!G{B[0]}:G{B[1]})", FMT_CURRENCY0, DUSTY_ROSE)
    kpi_card(ws, cards_row, 8, "Remaining Budget", f'=IF(TotalBudget="","",TotalBudget-SUM(Budget!G{B[0]}:G{B[1]}))', FMT_CURRENCY0, SAGE)
    kpi_card(ws, cards_row, 11, "Amount Paid", f"=SUM(Budget!J{B[0]}:J{B[1]})", FMT_CURRENCY0, SAGE)
    kpi_card(ws, cards_row, 14, "Outstanding", f"=SUM(Budget!K{B[0]}:K{B[1]})", FMT_CURRENCY0, GOLD)

    row = cards_row + 4
    sect = ws.cell(row=row, column=2, value="GUEST SNAPSHOT")
    sect.font = F_SECTION
    row += 1
    cards_row2 = row
    kpi_card(ws, cards_row2, 2, "Total Guests", f"=COUNTA('Guest List'!B{G[0]}:B{G[1]})", FMT_INT, DEEP_ROSE)
    kpi_card(ws, cards_row2, 5, "Confirmed", f"=COUNTIF('Guest List'!I{G[0]}:I{G[1]},\"Attending\")", FMT_INT, SAGE)
    kpi_card(ws, cards_row2, 8, "Pending RSVPs", f"=COUNTIF('Guest List'!I{G[0]}:I{G[1]},\"Pending\")", FMT_INT, GOLD)
    kpi_card(ws, cards_row2, 11, "Not Attending", f"=COUNTIF('Guest List'!I{G[0]}:I{G[1]},\"Declined\")", FMT_INT, "9D646B")
    kpi_card(ws, cards_row2, 14, "Vendors Booked", f"=COUNTA(Vendors!B{V[0]}:B{V[1]})", FMT_INT, DUSTY_ROSE)

    row = cards_row2 + 4
    sect = ws.cell(row=row, column=2, value="PLANNING PROGRESS & HONEYMOON")
    sect.font = F_SECTION
    row += 1
    cards_row3 = row
    kpi_card(ws, cards_row3, 2, "Checklist Complete", f"=IFERROR(COUNTIF('Master Checklist'!H{C[0]}:H{C[1]},\"Completed\")/(COUNTA('Master Checklist'!C{C[0]}:C{C[1]})),\"\")", FMT_PCT1, SAGE)
    kpi_card(ws, cards_row3, 5, "Tasks Remaining", f"=COUNTIFS('Master Checklist'!C{C[0]}:C{C[1]},\"<>\",'Master Checklist'!H{C[0]}:H{C[1]},\"<>Completed\")", FMT_INT, GOLD)
    kpi_card(ws, cards_row3, 8, "Outstanding Vendor Pmts", f"=SUM(Vendors!K{V[0]}:K{V[1]})", FMT_CURRENCY0, GOLD)
    kpi_card(ws, cards_row3, 11, "Honeymoon Budget", "=HoneymoonBudgetSetting", FMT_CURRENCY0, DEEP_ROSE)
    kpi_card(ws, cards_row3, 14, "Honeymoon Remaining", f"=SUM('Honeymoon Budget'!G{H[0]}:G{H[1]})", FMT_CURRENCY0, SAGE)

    # ---- Attention Needed ----
    row = cards_row3 + 4
    sect = ws.cell(row=row, column=2, value="ATTENTION NEEDED")
    sect.font = F_SECTION
    row += 1
    att_row = row
    kpi_card(ws, att_row, 2, "Overdue Payments",
             f'=COUNTIFS(Payments!H{P[0]}:H{P[1]},"<>",Payments!H{P[0]}:H{P[1]},"<"&TODAY(),Payments!I{P[0]}:I{P[1]},"<>Paid in Full")',
             FMT_INT, "9D646B")
    kpi_card(ws, att_row, 5, "Pending RSVPs",
             f'=COUNTIF(\'Guest List\'!I{G[0]}:I{G[1]},"Pending")', FMT_INT, GOLD)
    kpi_card(ws, att_row, 8, "Overdue Tasks",
             f'=COUNTIFS(\'Master Checklist\'!G{C[0]}:G{C[1]},"<>",\'Master Checklist\'!G{C[0]}:G{C[1]},"<"&TODAY(),\'Master Checklist\'!H{C[0]}:H{C[1]},"<>Completed")',
             FMT_INT, "9D646B")
    kpi_card(ws, att_row, 11, "Over-Budget Categories",
             "=SUMPRODUCT(($S$2:$S$21>$R$2:$R$21)*1)",
             FMT_INT, "9D646B")

    charts_start = att_row + 4

    # ================= HIDDEN CHART-DATA HELPER AREA =================
    BC_COL = get_column_letter(LIST_COL["BudgetCategory"])       # Lists col for budget categories
    RSVP_COL = get_column_letter(LIST_COL["RSVPStatus"])
    MEAL_COL = get_column_letter(LIST_COL["MealPreference"])
    PERIOD_COL = get_column_letter(LIST_COL["ChecklistPeriod"])

    ws.cell(row=1, column=18, value="Planned")
    ws.cell(row=1, column=19, value="Actual")
    for i in range(20):
        r = 2 + i
        ws.cell(row=r, column=18, value=f"=SUMIF(Budget!$B${B[0]}:$B${B[1]},Lists!{BC_COL}{r},Budget!$F${B[0]}:$F${B[1]})")  # R planned
        ws.cell(row=r, column=19, value=f"=SUMIF(Budget!$B${B[0]}:$B${B[1]},Lists!{BC_COL}{r},Budget!$G${B[0]}:$G${B[1]})")  # S actual
    ws.cell(row=2, column=21, value="Paid"); ws.cell(row=2, column=22, value=f"=SUM(Budget!J{B[0]}:J{B[1]})")
    ws.cell(row=3, column=21, value="Remaining"); ws.cell(row=3, column=22, value=f"=SUM(Budget!K{B[0]}:K{B[1]})")
    for k in range(1, 6):
        r = 1 + k
        ws.cell(row=r, column=23, value=f"=IFERROR(LARGE(Budget!$G${B[0]}:$G${B[1]},{k}),\"\")")
        ws.cell(row=r, column=24, value=f"=IFERROR(INDEX(Budget!$D${B[0]}:$D${B[1]},MATCH(W{r},Budget!$G${B[0]}:$G${B[1]},0)),\"\")")
    for i in range(4):
        r = 2 + i
        ws.cell(row=r, column=25, value=f"=COUNTIF('Guest List'!$I${G[0]}:$I${G[1]},Lists!{RSVP_COL}{r})")
    for i in range(6):
        r = 2 + i
        ws.cell(row=r, column=26, value=f"=COUNTIF('Guest List'!$L${G[0]}:$L${G[1]},Lists!{MEAL_COL}{r})")
    ws.cell(row=1, column=27, value="Completed")
    ws.cell(row=1, column=28, value="Remaining")
    for i in range(9):
        r = 2 + i
        ws.cell(row=r, column=27, value=f"=COUNTIFS('Master Checklist'!$B${C[0]}:$B${C[1]},Lists!{PERIOD_COL}{r},'Master Checklist'!$H${C[0]}:$H${C[1]},\"Completed\")")
        ws.cell(row=r, column=28, value=f"=COUNTIFS('Master Checklist'!$B${C[0]}:$B${C[1]},Lists!{PERIOD_COL}{r})-AA{r}")
    for col in ["R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB"]:
        ws.column_dimensions[col].hidden = True

    # ================= CHARTS =================
    catref = Reference(wb["Lists"], min_col=LIST_COL["BudgetCategory"], min_row=2, max_row=21)
    c1 = BarChart(); c1.type = "col"
    c1.add_data(Reference(ws, min_col=18, max_col=19, min_row=1, max_row=21), titles_from_data=True)
    c1.set_categories(catref)
    style_chart(c1, "Budget vs. Actual by Category")
    c1.y_axis.numFmt = FMT_CURRENCY0
    ws.add_chart(c1, f"B{charts_start}")

    c2 = DoughnutChart()
    c2.add_data(Reference(ws, min_col=19, min_row=2, max_row=21), titles_from_data=False)
    c2.set_categories(catref)
    style_chart(c2, "Expenses by Category")
    colorize_doughnut(c2, 20)
    c2.dataLabels = DataLabelList(); c2.dataLabels.showPercent = True
    ws.add_chart(c2, f"H{charts_start}")

    c3 = BarChart(); c3.type = "bar"
    c3.add_data(Reference(ws, min_col=22, min_row=2, max_row=3), titles_from_data=False)
    c3.set_categories(Reference(ws, min_col=21, min_row=2, max_row=3))
    style_chart(c3, "Paid vs. Remaining", height=6.5)
    c3.legend = None
    ws.add_chart(c3, f"B{charts_start+17}")

    c4 = BarChart(); c4.type = "bar"
    c4.add_data(Reference(ws, min_col=23, min_row=2, max_row=6), titles_from_data=False)
    c4.set_categories(Reference(ws, min_col=24, min_row=2, max_row=6))
    style_chart(c4, "Top 5 Wedding Expenses", height=6.5)
    c4.legend = None
    ws.add_chart(c4, f"H{charts_start+17}")

    c5 = DoughnutChart()
    c5.add_data(Reference(ws, min_col=25, min_row=2, max_row=5), titles_from_data=False)
    c5.set_categories(Reference(wb["Lists"], min_col=LIST_COL["RSVPStatus"], min_row=2, max_row=5))
    style_chart(c5, "RSVP Status", height=8)
    colorize_doughnut(c5, 4)
    c5.dataLabels = DataLabelList(); c5.dataLabels.showVal = True
    ws.add_chart(c5, f"B{charts_start+31}")

    c6 = BarChart(); c6.type = "col"
    c6.add_data(Reference(ws, min_col=26, min_row=2, max_row=7), titles_from_data=False)
    c6.set_categories(Reference(wb["Lists"], min_col=LIST_COL["MealPreference"], min_row=2, max_row=7))
    style_chart(c6, "Guest Meal Preferences", height=8)
    c6.legend = None
    ws.add_chart(c6, f"H{charts_start+31}")

    c7 = BarChart(); c7.type = "col"; c7.grouping = "stacked"; c7.overlap = 100
    c7.add_data(Reference(ws, min_col=27, min_row=1, max_row=10), titles_from_data=True)
    c7.add_data(Reference(ws, min_col=28, min_row=1, max_row=10), titles_from_data=True)
    c7.set_categories(Reference(wb["Lists"], min_col=LIST_COL["ChecklistPeriod"], min_row=2, max_row=10))
    style_chart(c7, "Checklist Completion by Period", height=8.5, width=15.5)
    ws.add_chart(c7, f"B{charts_start+45}")

    c8 = BarChart(); c8.type = "col"
    c8.add_data(Reference(wb["Honeymoon Budget"], min_col=3, min_row=H[0]-1, max_row=H[1]), titles_from_data=True)
    c8.add_data(Reference(wb["Honeymoon Budget"], min_col=4, min_row=H[0]-1, max_row=H[1]), titles_from_data=True)
    c8.set_categories(Reference(wb["Honeymoon Budget"], min_col=2, min_row=H[0], max_row=H[1]))
    style_chart(c8, "Honeymoon: Planned vs. Actual", height=8.5, width=15.5)
    ws.add_chart(c8, f"H{charts_start+45}")

    for col, w in zip("ABCDEFGHIJKLMNOP", [3]*16):
        ws.column_dimensions[col].width = 8.2
    ws.page_setup.orientation = "landscape"
    ws.sheet_view.zoomScale = 85
    return ws
