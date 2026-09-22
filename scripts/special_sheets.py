from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from styles import (
    sheet_header, set_sheet_defaults, section_title, F_BODY, F_HEADER_LIGHT, FILL_HEADER,
    ALIGN_LEFT, ALIGN_LEFT_WRAP, ALIGN_CENTER, FMT_CURRENCY0, FMT_DATE, FMT_INT, F_SECTION,
    F_KPI_LABEL, kpi_card, col_idx, DEEP_ROSE, GOLD, FILL_CREAM, B_ALL_LIGHT, FILL_LIGHT_SAGE,
    fill, LIGHT_SAGE, LIGHT_BLUSH, F_NAV,
)
from tablesheet import status_conditional_formatting
from openpyxl.formatting.rule import FormulaRule

NAV_HOME = [("Dashboard", "Dashboard"), ("Start Here", "START HERE")]


# ------------------------------------------------------------ SEATING PLAN
def build_seating(wb, guest_first, guest_last, demo):
    ws = wb.create_sheet("Seating Plan")
    set_sheet_defaults(ws)
    row = sheet_header(ws, "SEATING PLAN", "Assign every guest to a table — synced automatically with your Guest List.", "K", NAV_HOME)

    row = section_title(ws, row, 2, "Table Overview", "K")
    headers = ["Table #", "Capacity", "Seated", "Seats Remaining"]
    # Two side-by-side blocks (tables 1-10, 11-20) so this reference section stays
    # compact — keeping the freeze pane below it from swallowing too much of the screen.
    BLOCK_COLS = [2, 8]  # B..E and H..K
    for block_col in BLOCK_COLS:
        for i, h in enumerate(headers):
            c = ws.cell(row=row, column=block_col + i, value=h)
            c.font = F_HEADER_LIGHT
            c.fill = FILL_HEADER
            c.alignment = ALIGN_CENTER
    tbl_row0 = row + 1
    n_tables = 20
    half = 10
    for t in range(1, n_tables + 1):
        block_col = BLOCK_COLS[0] if t <= half else BLOCK_COLS[1]
        r = tbl_row0 + ((t - 1) % half)
        b, cpc, sc, rc = block_col, block_col + 1, block_col + 2, block_col + 3
        ws.cell(row=r, column=b, value=t).alignment = ALIGN_CENTER
        cap = ws.cell(row=r, column=cpc, value="=IF({bl}{r}<=TableCount,SeatsPerTable,\"\")".format(bl=get_column_letter(b), r=r))
        cap.number_format = FMT_INT
        cap.alignment = ALIGN_CENTER
        seated = ws.cell(row=r, column=sc,
                          value=f"=IF({get_column_letter(cpc)}{r}=\"\",\"\",COUNTIF(GuestTable[Table Number],{get_column_letter(b)}{r}))")
        seated.number_format = FMT_INT
        seated.alignment = ALIGN_CENTER
        remaining = ws.cell(row=r, column=rc,
                             value=f'=IF({get_column_letter(cpc)}{r}="","",{get_column_letter(cpc)}{r}-{get_column_letter(sc)}{r})')
        remaining.number_format = FMT_INT
        remaining.alignment = ALIGN_CENTER
        for cc in range(b, rc + 1):
            ws.cell(row=r, column=cc).border = B_ALL_LIGHT
            ws.cell(row=r, column=cc).font = F_BODY
            if r % 2 == 0:
                ws.cell(row=r, column=cc).fill = fill("FBF5F0")
    for block_col in BLOCK_COLS:
        rc_letter = get_column_letter(block_col + 3)
        ws.conditional_formatting.add(f"{rc_letter}{tbl_row0}:{rc_letter}{tbl_row0 + half - 1}", FormulaRule(
            formula=[f'AND({rc_letter}{tbl_row0}<>"",{rc_letter}{tbl_row0}<0)'], fill=fill("F3D9D6")))

    roster_row = tbl_row0 + half + 2
    roster_row = section_title(ws, roster_row, 2, "Guest Roster", "K")
    headers2 = ["Table Number", "Guest Name", "Household", "Meal", "Special Requirements", "Notes"]
    for i, h in enumerate(headers2):
        c = ws.cell(row=roster_row, column=2 + i, value=h)
        c.font = F_HEADER_LIGHT
        c.fill = FILL_HEADER
        c.alignment = ALIGN_CENTER
    n_guests = guest_last - guest_first + 1
    for i in range(n_guests):
        grow = guest_first + i
        r = roster_row + 1 + i
        ws.cell(row=r, column=2, value=f"=IF('Guest List'!O{grow}=\"\",\"\",'Guest List'!O{grow})").alignment = ALIGN_CENTER
        ws.cell(row=r, column=3, value=f"=IF('Guest List'!B{grow}=\"\",\"\",'Guest List'!B{grow})")
        ws.cell(row=r, column=4, value=f"=IF('Guest List'!C{grow}=\"\",\"\",'Guest List'!C{grow})")
        ws.cell(row=r, column=5, value=f"=IF('Guest List'!L{grow}=\"\",\"\",'Guest List'!L{grow})").alignment = ALIGN_CENTER
        for cc in range(2, 8):
            cell = ws.cell(row=r, column=cc)
            cell.font = F_BODY
            cell.border = B_ALL_LIGHT
            if cell.value is None:
                cell.value = None
            if r % 2 == 0:
                cell.fill = fill("FBF5F0")
    for i, w in enumerate([13, 22, 20, 15, 24, 22]):
        ws.column_dimensions[get_column_letter(2 + i)].width = w
    for c in BLOCK_COLS:
        for j in range(4):
            ws.column_dimensions[get_column_letter(c + j)].width = max(
                ws.column_dimensions[get_column_letter(c + j)].width or 0, 13)
    # Freeze just below the roster's own header row — the Table Overview above it
    # is now a compact two-column block rather than a 20-row single column, so far
    # less of the screen is consumed by the frozen region than before.
    ws.freeze_panes = f"B{roster_row+1}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:K{roster_row + n_guests}"
    ws.print_title_rows = f"{roster_row}:{roster_row}"
    return ws


# ---------------------------------------------------------------- CALENDAR
def build_calendar(wb, checklist_first, checklist_last, payments_first, payments_last, demo):
    ws = wb.create_sheet("Calendar")
    set_sheet_defaults(ws)
    row = sheet_header(ws, "PLANNING CALENDAR", "Every task deadline and payment due date, gathered in one chronological view.", "E", NAV_HOME)

    row = section_title(ws, row, 2, "Task Deadlines", "E")
    hdrs = ["Due Date", "Task", "Category", "Status"]
    for i, h in enumerate(hdrs):
        c = ws.cell(row=row, column=2 + i, value=h)
        c.font = F_HEADER_LIGHT
        c.fill = FILL_HEADER
        c.alignment = ALIGN_CENTER
    r0 = row + 1
    n = checklist_last - checklist_first + 1
    for i in range(n):
        crow = checklist_first + i
        r = r0 + i
        d = ws.cell(row=r, column=2, value=f"=IF('Master Checklist'!G{crow}=\"\",\"\",'Master Checklist'!G{crow})")
        d.number_format = FMT_DATE
        d.alignment = ALIGN_CENTER
        ws.cell(row=r, column=3, value=f"=IF('Master Checklist'!C{crow}=\"\",\"\",'Master Checklist'!C{crow})")
        ws.cell(row=r, column=4, value=f"=IF('Master Checklist'!D{crow}=\"\",\"\",'Master Checklist'!D{crow})")
        st = ws.cell(row=r, column=5, value=f"=IF('Master Checklist'!H{crow}=\"\",\"\",'Master Checklist'!H{crow})")
        st.alignment = ALIGN_CENTER
        for cc in range(2, 6):
            cell = ws.cell(row=r, column=cc)
            cell.font = F_BODY
            cell.border = B_ALL_LIGHT
            if r % 2 == 0:
                cell.fill = fill("FBF5F0")
    status_conditional_formatting(ws, f"E{r0}:E{r0+n-1}", {"Completed": LIGHT_SAGE, "Overdue": "F3D9D6", "In Progress": LIGHT_BLUSH})
    ws.conditional_formatting.add(f"B{r0}:B{r0+n-1}", FormulaRule(
        formula=[f'AND(B{r0}<>"",B{r0}<TODAY(),E{r0}<>"Completed")'], fill=fill("F3D9D6")))

    row2 = r0 + n + 2
    row2 = section_title(ws, row2, 2, "Payment Deadlines", "E")
    for i, h in enumerate(["Due Date", "Vendor / Expense", "Amount", "Status"]):
        c = ws.cell(row=row2, column=2 + i, value=h)
        c.font = F_HEADER_LIGHT
        c.fill = FILL_HEADER
        c.alignment = ALIGN_CENTER
    r1 = row2 + 1
    m = payments_last - payments_first + 1
    for i in range(m):
        prow = payments_first + i
        r = r1 + i
        d = ws.cell(row=r, column=2, value=f"=IF(Payments!H{prow}=\"\",\"\",Payments!H{prow})")
        d.number_format = FMT_DATE
        d.alignment = ALIGN_CENTER
        ws.cell(row=r, column=3, value=f"=IF(Payments!B{prow}=\"\",\"\",Payments!B{prow}&\" - \"&Payments!C{prow})")
        amt = ws.cell(row=r, column=4, value=f"=IF(Payments!F{prow}=\"\",\"\",Payments!F{prow})")
        amt.number_format = FMT_CURRENCY0
        amt.alignment = ALIGN_LEFT
        st = ws.cell(row=r, column=5, value=f"=IF(Payments!I{prow}=\"\",\"\",Payments!I{prow})")
        st.alignment = ALIGN_CENTER
        for cc in range(2, 6):
            cell = ws.cell(row=r, column=cc)
            cell.font = F_BODY
            cell.border = B_ALL_LIGHT
            if r % 2 == 0:
                cell.fill = fill("FBF5F0")
    status_conditional_formatting(ws, f"E{r1}:E{r1+m-1}", {"Paid in Full": LIGHT_SAGE, "Overdue": "F3D9D6", "Deposit Paid": LIGHT_BLUSH})

    for i, w in enumerate([14, 30, 14, 16]):
        ws.column_dimensions[get_column_letter(2 + i)].width = w
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6
    ws.print_area = f"A1:E{r1 + m - 1}"
    ws.print_title_rows = f"{row}:{row},{row2}:{row2}"
    return ws


# ---------------------------------------------------------------- START HERE
def build_start_here(wb, demo):
    ws = wb.create_sheet("START HERE", 0)
    set_sheet_defaults(ws, GOLD)
    ws.sheet_view.showGridLines = False
    for r in range(1, 5):
        for c in range(1, 9):
            ws.cell(row=r, column=c).fill = FILL_LIGHT_SAGE
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 34
    t = ws.cell(row=2, column=2, value="WELCOME TO YOUR PREMIUM WEDDING PLANNER")
    from styles import F_SHEET_TITLE, F_SHEET_SUB
    t.font = F_SHEET_TITLE
    s = ws.cell(row=3, column=2, value="A calm, elegant system for planning every detail of your wedding — from budget to seating chart.")
    s.font = F_SHEET_SUB
    for c in range(1, 9):
        ws.cell(row=4, column=c).border = None
    from openpyxl.styles import Border, Side
    for c in range(1, 9):
        ws.cell(row=4, column=c).border = Border(bottom=Side(style="medium", color=GOLD))

    row = 6
    row = section_title(ws, row, 2, "How This Workbook Works", "H")
    tips = [
        "1. Start on the Settings sheet — enter your names, wedding date, budget and guest target once.",
        "2. Every other sheet (Dashboard, Budget, Guest List, Checklist and more) calculates automatically from what you enter.",
        "3. Use the dropdown menus in colored-header columns to keep your data consistent — just click a cell and choose from the list.",
        "4. Visit the Dashboard any time for a beautiful, real-time snapshot of your whole wedding.",
        "5. Each sheet has a small navigation link at the top-left to jump back to the Dashboard or this page.",
        "6. All key sheets are print-ready — use File > Print, and each will fit neatly on the page.",
    ]
    for tip in tips:
        c = ws.cell(row=row, column=2, value=tip)
        c.font = F_BODY
        c.alignment = ALIGN_LEFT_WRAP
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.row_dimensions[row].height = 18
        row += 1

    row += 1
    row = section_title(ws, row, 2, "What's Inside", "H")
    directory = [
        ("Dashboard", "Your wedding at a glance — budget, guests, checklist and more."),
        ("Budget", "Track every planned and actual expense by category."),
        ("Vendors", "Contact details, quotes and contracts for every vendor."),
        ("Guest List", "Manage invitations, RSVPs and meal preferences."),
        ("Seating Plan", "Assign guests to tables, synced with your Guest List."),
        ("Timeline", "Minute-by-minute plan for the wedding day."),
        ("Master Checklist", "Every planning task, organized by timeframe."),
        ("Payments", "Every deposit, balance and due date."),
        ("Wedding Party", "Roles and responsibilities for your closest circle."),
        ("Venue Comparison", "Compare every venue you're considering."),
        ("Food & Drinks", "Plan your menu and catering costs."),
        ("Honeymoon Budget", "Plan and track your post-wedding getaway."),
        ("Contacts", "Every phone number and email in one place."),
        ("Gift Tracker", "Track gifts received and thank-you notes."),
        ("Calendar", "A chronological view of tasks and payment deadlines."),
        ("Settings", "Your wedding details, budget and dropdown lists."),
    ]
    for name, desc in directory:
        c1 = ws.cell(row=row, column=2, value=name)
        c1.font = F_NAV
        c1.hyperlink = f"#'{name}'!A1"
        c2 = ws.cell(row=row, column=4, value=desc)
        c2.font = F_BODY
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=8)
        ws.row_dimensions[row].height = 16
        row += 1

    ws.column_dimensions["A"].width = 2
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 13
    ws.column_dimensions["B"].width = 18
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6
    ws.print_area = f"A1:H{row}"
    return ws
