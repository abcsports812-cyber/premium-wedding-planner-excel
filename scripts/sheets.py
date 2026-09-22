from openpyxl.utils import get_column_letter
from tablesheet import Col, build_table_sheet, status_conditional_formatting
from styles import (
    FMT_CURRENCY0, FMT_DATE, FMT_INT, FMT_PCT, FMT_PCT1, F_BODY, F_KPI_LABEL, F_SECTION,
    ALIGN_LEFT, DEEP_ROSE, SAGE, GOLD, LIGHT_SAGE, LIGHT_BLUSH, MUTED_GOLD, DUSTY_ROSE,
    CHARCOAL, fill, kpi_card, col_idx,
)
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
import demo_data as D

NAV_HOME = [("Dashboard", "Dashboard"), ("Start Here", "START HERE")]


def add_summary_strip(ws, row, last_col_letter, items):
    """items: list of (label, formula_str_no_eq, number_format)"""
    col = 2
    for label, formula, fmt in items:
        kpi_card(ws, row, col, label, "=" + formula, fmt, accent=DEEP_ROSE, height=3, width=3)
        col += 3
    return row + 4


# ---------------------------------------------------------------- BUDGET ---
def build_budget(wb, ranges, demo):
    cols = [
        Col("Category", 16, validation_list=ranges["BudgetCategory"]),
        Col("Subcategory", 20),
        Col("Item / Description", 30, wrap=True),
        Col("Vendor", 22),
        Col("Planned Budget", 15, FMT_CURRENCY0, align="right"),
        Col("Actual Cost", 15, FMT_CURRENCY0, align="right"),
        Col("Difference", 15, FMT_CURRENCY0, align="right",
            formula=lambda r: f'IF(F{r}="","",F{r}-G{r})'),
        Col("Deposit", 13, FMT_CURRENCY0, align="right"),
        Col("Amount Paid", 15, FMT_CURRENCY0, align="right"),
        Col("Amount Remaining", 17, FMT_CURRENCY0, align="right",
            formula=lambda r: f'IF(G{r}="","",G{r}-J{r})'),
        Col("Due Date", 14, FMT_DATE, align="center"),
        Col("Payment Status", 16, validation_list=ranges["PaymentStatus"], align="center"),
        Col("Notes", 24, wrap=True),
    ]
    demo_rows = []
    for cat, sub, item, vendor, planned, actual, deposit, paid, due in D.BUDGET_ITEMS:
        demo_rows.append({"Category": cat, "Subcategory": sub, "Item / Description": item,
                           "Vendor": vendor, "Planned Budget": planned, "Actual Cost": actual,
                           "Deposit": deposit, "Amount Paid": paid, "Due Date": due,
                           "Payment Status": "Paid in Full" if paid >= actual else ("Deposit Paid" if paid > 0 else "Not Started"),
                           "Notes": ""})
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Budget", "BUDGET & EXPENSES", "Track your wedding spending and stay within your plan.",
        cols, demo_rows, "BudgetTable", 30, demo, nav_links=NAV_HOME)

    status_conditional_formatting(ws, f"L{first}:L{last}", {
        "Paid in Full": LIGHT_SAGE, "Overdue": "F3D9D6", "Deposit Paid": LIGHT_BLUSH,
        "Not Started": MUTED_GOLD,
    })
    ws.conditional_formatting.add(f"G{first}:G{last}", CellIsRule(
        operator="greaterThan", formula=[f"F{first}"], fill=fill("F3D9D6")))

    srow = last + 2
    ws.cell(row=srow, column=2, value="BUDGET SUMMARY").font = F_SECTION
    srow += 1
    add_summary_strip(ws, srow, lastcol, [
        ("Total Planned", f"SUM(F{first}:F{last})", FMT_CURRENCY0),
        ("Total Actual", f"SUM(G{first}:G{last})", FMT_CURRENCY0),
        ("Total Remaining", f"SUM(K{first}:K{last})", FMT_CURRENCY0),
        ("Budget Used", f'IFERROR(SUM(G{first}:G{last})/TotalBudget,"")', FMT_PCT1),
    ])
    return ws, first, last


# --------------------------------------------------------------- VENDORS ---
def build_vendors(wb, ranges, demo):
    cols = [
        Col("Vendor Name", 24),
        Col("Category", 16, validation_list=ranges["BudgetCategory"]),
        Col("Contact Person", 18),
        Col("Phone", 15),
        Col("Email", 24),
        Col("Website", 20),
        Col("Quote", 13, FMT_CURRENCY0, align="right"),
        Col("Final Cost", 13, FMT_CURRENCY0, align="right"),
        Col("Deposit", 13, FMT_CURRENCY0, align="right"),
        Col("Remaining", 13, FMT_CURRENCY0, align="right",
            formula=lambda r: f'IF(I{r}="","",I{r}-J{r})'),
        Col("Contract Status", 15, validation_list=ranges["ContractStatus"], align="center"),
        Col("Payment Status", 15, validation_list=ranges["PaymentStatus"], align="center"),
        Col("Due Date", 13, FMT_DATE, align="center"),
        Col("Rating", 9, validation_list=ranges["Rating"], align="center"),
        Col("Notes", 22, wrap=True),
    ]
    demo_rows = []
    for name, cat, contact, phone, email, site, quote, final, dep, contract, pay, due, rating in D.VENDORS:
        demo_rows.append({"Vendor Name": name, "Category": cat, "Contact Person": contact,
                           "Phone": phone, "Email": email, "Website": site, "Quote": quote,
                           "Final Cost": final, "Deposit": dep, "Contract Status": contract,
                           "Payment Status": pay, "Due Date": due, "Rating": rating, "Notes": ""})
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Vendors", "VENDOR DIRECTORY", "Every vendor, quote and contract in one elegant place.",
        cols, demo_rows, "VendorsTable", 26, demo, nav_links=NAV_HOME)
    status_conditional_formatting(ws, f"L{first}:L{last}", {"Signed": LIGHT_SAGE, "Sent": MUTED_GOLD, "Not Sent": "F3D9D6"})
    status_conditional_formatting(ws, f"M{first}:M{last}", {
        "Paid in Full": LIGHT_SAGE, "Overdue": "F3D9D6", "Deposit Paid": LIGHT_BLUSH, "Not Started": MUTED_GOLD})
    srow = last + 2
    ws.cell(row=srow, column=2, value="VENDOR SUMMARY").font = F_SECTION
    srow += 1
    add_summary_strip(ws, srow, lastcol, [
        ("Total Vendors", f'COUNTA(B{first}:B{last})', FMT_INT),
        ("Contracts Signed", f'COUNTIF(L{first}:L{last},"Signed")', FMT_INT),
        ("Outstanding Balance", f"SUM(K{first}:K{last})", FMT_CURRENCY0),
        ("Overdue Payments", f'COUNTIF(M{first}:M{last},"Overdue")', FMT_INT),
    ])
    return ws, first, last


# ------------------------------------------------------------ GUEST LIST ---
def build_guest_list(wb, ranges, demo):
    cols = [
        Col("Guest Name", 22),
        Col("Household / Group", 20),
        Col("Side", 12, validation_list=ranges["Side"], align="center"),
        Col("Relationship", 18),
        Col("Email", 24),
        Col("Phone", 15),
        Col("Invitation Sent", 16, validation_list=ranges["InvitationStatus"], align="center"),
        Col("RSVP Status", 14, validation_list=ranges["RSVPStatus"], align="center"),
        Col("Plus One", 10, validation_list=ranges["PlusOne"], align="center"),
        Col("Plus One Name", 20),
        Col("Meal Preference", 16, validation_list=ranges["MealPreference"], align="center"),
        Col("Dietary Restrictions", 18),
        Col("Children Attending", 12, validation_list=ranges["Children"], align="center"),
        Col("Table Number", 12, align="center"),
        Col("Notes", 20, wrap=True),
    ]
    demo_rows = D.build_guest_rows(58) if demo else []
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Guest List", "GUEST LIST & RSVP", "Manage every invitee, RSVP and meal preference.",
        cols, demo_rows, "GuestTable", 60, demo, nav_links=NAV_HOME)
    status_conditional_formatting(ws, f"I{first}:I{last}", {
        "Attending": LIGHT_SAGE, "Declined": "F3D9D6", "Pending": MUTED_GOLD, "Tentative": LIGHT_BLUSH})
    srow = last + 2
    ws.cell(row=srow, column=2, value="GUEST SUMMARY").font = F_SECTION
    srow += 1
    add_summary_strip(ws, srow, lastcol, [
        ("Total Invited", f'COUNTA(B{first}:B{last})', FMT_INT),
        ("Attending", f'COUNTIF(I{first}:I{last},"Attending")+COUNTIFS(I{first}:I{last},"Attending",J{first}:J{last},"Yes")', FMT_INT),
        ("Pending", f'COUNTIF(I{first}:I{last},"Pending")', FMT_INT),
        ("Declined", f'COUNTIF(I{first}:I{last},"Declined")', FMT_INT),
    ])
    return ws, first, last


# --------------------------------------------------------------- PAYMENTS --
def build_payments(wb, ranges, demo):
    cols = [
        Col("Vendor", 22, validation_list=None),
        Col("Expense", 26, wrap=True),
        Col("Total Cost", 14, FMT_CURRENCY0, align="right"),
        Col("Deposit", 13, FMT_CURRENCY0, align="right"),
        Col("Amount Paid", 14, FMT_CURRENCY0, align="right"),
        Col("Remaining", 14, FMT_CURRENCY0, align="right",
            formula=lambda r: f'IF(D{r}="","",D{r}-(E{r}+F{r}))'),
        Col("Due Date", 13, FMT_DATE, align="center"),
        Col("Payment Status", 15, validation_list=ranges["PaymentStatus"], align="center"),
        Col("Payment Method", 16, validation_list=ranges["PaymentMethod"], align="center"),
        Col("Notes", 20, wrap=True),
    ]
    demo_rows = []
    for vendor, expense, total, dep, paid, remaining, due, status, method in D.PAYMENTS:
        demo_rows.append({"Vendor": vendor, "Expense": expense, "Total Cost": total,
                           "Deposit": dep, "Amount Paid": paid, "Due Date": due,
                           "Payment Status": status, "Payment Method": method, "Notes": ""})
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Payments", "PAYMENT SCHEDULE", "Every deposit, balance and due date, tracked automatically.",
        cols, demo_rows, "PaymentsTable", 24, demo, nav_links=NAV_HOME)
    status_conditional_formatting(ws, f"I{first}:I{last}", {
        "Paid in Full": LIGHT_SAGE, "Overdue": "F3D9D6", "Deposit Paid": LIGHT_BLUSH,
        "Partially Paid": LIGHT_BLUSH, "Not Started": MUTED_GOLD})
    ws.conditional_formatting.add(f"H{first}:H{last}", FormulaRule(
        formula=[f'AND(H{first}<TODAY(),I{first}<>"Paid in Full",H{first}<>"")'], fill=fill("F3D9D6")))
    srow = last + 2
    ws.cell(row=srow, column=2, value="PAYMENTS SUMMARY").font = F_SECTION
    srow += 1
    add_summary_strip(ws, srow, lastcol, [
        ("Total Payable", f"SUM(D{first}:D{last})", FMT_CURRENCY0),
        ("Total Paid", f"SUM(E{first}:E{last})+SUM(F{first}:F{last})", FMT_CURRENCY0),
        ("Total Remaining", f"SUM(G{first}:G{last})", FMT_CURRENCY0),
        ("Overdue", f'COUNTIFS(H{first}:H{last},"<"&TODAY(),I{first}:I{last},"<>Paid in Full")', FMT_INT),
    ])
    return ws, first, last


# ---------------------------------------------------------------- CHECKLIST
def build_checklist(wb, ranges, demo):
    cols = [
        Col("Period", 18, validation_list=ranges["ChecklistPeriod"]),
        Col("Task", 32, wrap=True),
        Col("Category", 16),
        Col("Priority", 11, validation_list=ranges["Priority"], align="center"),
        Col("Owner", 16, validation_list=ranges["Owner"], align="center"),
        Col("Due Date", 13, FMT_DATE, align="center"),
        Col("Status", 14, validation_list=ranges["TaskStatus"], align="center"),
        Col("Notes", 22, wrap=True),
    ]
    demo_rows = []
    for period, task, cat, pri, owner, due, status in D.CHECKLIST_TASKS:
        demo_rows.append({"Period": period, "Task": task, "Category": cat, "Priority": pri,
                           "Owner": owner, "Due Date": due, "Status": status, "Notes": ""})
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Master Checklist", "MASTER CHECKLIST", "Every planning task, organized by timeline.",
        cols, demo_rows, "ChecklistTable", 60, demo, nav_links=NAV_HOME)
    status_conditional_formatting(ws, f"H{first}:H{last}", {
        "Completed": LIGHT_SAGE, "In Progress": LIGHT_BLUSH, "Not Started": "FFFFFF", "Overdue": "F3D9D6"})
    ws.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(
        formula=[f'AND(G{first}<TODAY(),H{first}<>"Completed",G{first}<>"")'], fill=fill("F3D9D6")))
    ws.conditional_formatting.add(f"E{first}:E{last}", FormulaRule(
        formula=[f'EXACT(E{first},"High")'], font=None, fill=fill(MUTED_GOLD)))
    srow = last + 2
    ws.cell(row=srow, column=2, value="CHECKLIST SUMMARY").font = F_SECTION
    srow += 1
    add_summary_strip(ws, srow, lastcol, [
        ("Total Tasks", f'COUNTA(C{first}:C{last})', FMT_INT),
        ("Completed", f'COUNTIF(H{first}:H{last},"Completed")', FMT_INT),
        ("Remaining", f'COUNTIFS(C{first}:C{last},"<>",H{first}:H{last},"<>Completed")', FMT_INT),
        ("Completion %", f'IFERROR(COUNTIF(H{first}:H{last},"Completed")/(COUNTA(C{first}:C{last})),"")', FMT_PCT1),
    ])
    return ws, first, last


# ------------------------------------------------------------ WEDDING PARTY
def build_wedding_party(wb, ranges, demo):
    cols = [
        Col("Name", 20), Col("Role", 16), Col("Side", 12, validation_list=ranges["Side"], align="center"),
        Col("Phone", 15), Col("Email", 24), Col("Responsibilities", 28, wrap=True),
        Col("Attire", 18), Col("Notes", 16, wrap=True),
    ]
    demo_rows = []
    for name, role, side, phone, email, resp, attire, notes in D.WEDDING_PARTY:
        demo_rows.append({"Name": name, "Role": role, "Side": side, "Phone": phone, "Email": email,
                           "Responsibilities": resp, "Attire": attire, "Notes": notes})
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Wedding Party", "WEDDING PARTY", "Roles, responsibilities and attire for your closest circle.",
        cols, demo_rows, "WeddingPartyTable", 12, demo, nav_links=NAV_HOME)
    return ws, first, last


# ---------------------------------------------------------- VENUE COMPARISON
def build_venue_comparison(wb, ranges, demo):
    cols = [
        Col("Venue", 20), Col("Location", 18), Col("Capacity", 10, FMT_INT, align="center"),
        Col("Rental Cost", 13, FMT_CURRENCY0, align="right"), Col("Catering Cost", 13, FMT_CURRENCY0, align="right"),
        Col("Deposit", 11, FMT_CURRENCY0, align="right"), Col("Availability", 18, align="center"),
        Col("Ceremony Included", 12, validation_list=ranges["YesNo"], align="center"),
        Col("Reception Included", 12, validation_list=ranges["YesNo"], align="center"),
        Col("Parking", 10, validation_list=ranges["YesNo"], align="center"),
        Col("Accommodation", 12, validation_list=ranges["YesNo"], align="center"),
        Col("Rating", 9, validation_list=ranges["Rating"], align="center"), Col("Notes", 24, wrap=True),
    ]
    demo_rows = []
    for v, loc, cap, rc, cc, dep, avail, cer, rec, park, acc, rating, notes in D.VENUES:
        demo_rows.append({"Venue": v, "Location": loc, "Capacity": cap, "Rental Cost": rc,
                           "Catering Cost": cc, "Deposit": dep, "Availability": avail,
                           "Ceremony Included": cer, "Reception Included": rec, "Parking": park,
                           "Accommodation": acc, "Rating": rating, "Notes": notes})
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Venue Comparison", "VENUE COMPARISON", "Compare every venue you toured, side by side.",
        cols, demo_rows, "VenueTable", 10, demo, nav_links=NAV_HOME)
    return ws, first, last


# -------------------------------------------------------------- FOOD&DRINKS
def build_food_drinks(wb, ranges, demo):
    cols = [
        Col("Catering Vendor", 22), Col("Menu Item", 28, wrap=True),
        Col("Category", 16, validation_list=ranges["FoodCategory"], align="center"),
        Col("Quantity", 10, FMT_INT, align="center"), Col("Price", 11, FMT_CURRENCY0, align="right"),
        Col("Total", 13, FMT_CURRENCY0, align="right", formula=lambda r: f'IF(E{r}="","",E{r}*F{r})'),
        Col("Dietary Notes", 18), Col("Notes", 20, wrap=True),
    ]
    demo_rows = []
    for vendor, item, cat, qty, price, diet, notes in D.FOOD_DRINKS:
        demo_rows.append({"Catering Vendor": vendor, "Menu Item": item, "Category": cat,
                           "Quantity": qty, "Price": price, "Dietary Notes": diet, "Notes": notes})
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Food & Drinks", "FOOD & DRINKS", "Plan your menu, portions and per-guest cost.",
        cols, demo_rows, "FoodTable", 24, demo, nav_links=NAV_HOME)
    srow = last + 2
    ws.cell(row=srow, column=2, value="MENU SUMMARY").font = F_SECTION
    srow += 1
    add_summary_strip(ws, srow, lastcol, [
        ("Total Menu Cost", f"SUM(G{first}:G{last})", FMT_CURRENCY0),
    ])
    return ws, first, last


# ----------------------------------------------------------------- HONEYMOON
def build_honeymoon(wb, ranges, demo):
    cols = [
        Col("Category", 18, validation_list=ranges["HoneymoonCategory"]),
        Col("Planned Budget", 15, FMT_CURRENCY0, align="right"),
        Col("Actual Cost", 14, FMT_CURRENCY0, align="right"),
        Col("Difference", 14, FMT_CURRENCY0, align="right", formula=lambda r: f'IF(C{r}="","",C{r}-D{r})'),
        Col("Amount Paid", 14, FMT_CURRENCY0, align="right"),
        Col("Remaining", 13, FMT_CURRENCY0, align="right", formula=lambda r: f'IF(D{r}="","",D{r}-F{r})'),
        Col("Due Date", 13, FMT_DATE, align="center"), Col("Notes", 22, wrap=True),
    ]
    demo_rows = []
    for cat, planned, actual, paid, remaining, due in D.HONEYMOON:
        demo_rows.append({"Category": cat, "Planned Budget": planned, "Actual Cost": actual,
                           "Amount Paid": paid, "Due Date": due, "Notes": ""})
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Honeymoon Budget", "HONEYMOON BUDGET", "Plan and track spending for your post-wedding getaway.",
        cols, demo_rows, "HoneymoonTable", 10, demo, nav_links=NAV_HOME)
    srow = last + 2
    ws.cell(row=srow, column=2, value="HONEYMOON SUMMARY").font = F_SECTION
    srow += 1
    add_summary_strip(ws, srow, lastcol, [
        ("Planned Total", f"SUM(C{first}:C{last})", FMT_CURRENCY0),
        ("Actual Total", f"SUM(D{first}:D{last})", FMT_CURRENCY0),
        ("Remaining", f"SUM(G{first}:G{last})", FMT_CURRENCY0),
    ])
    return ws, first, last


# ----------------------------------------------------------------- CONTACTS
def build_contacts(wb, ranges, demo):
    cols = [
        Col("Name", 20), Col("Company", 22), Col("Category", 16, validation_list=ranges["BudgetCategory"]),
        Col("Role", 18), Col("Phone", 15), Col("Email", 24), Col("Website", 18),
        Col("Address", 26, wrap=True), Col("Notes", 20, wrap=True),
    ]
    demo_rows = []
    for name, company, cat, role, phone, email, site, addr, notes in D.CONTACTS:
        demo_rows.append({"Name": name, "Company": company, "Category": cat, "Role": role,
                           "Phone": phone, "Email": email, "Website": site, "Address": addr, "Notes": notes})
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Contacts", "CONTACTS", "Every phone number and email you need, in one place.",
        cols, demo_rows, "ContactsTable", 20, demo, nav_links=NAV_HOME)
    return ws, first, last


# --------------------------------------------------------------- TIMELINE --
def build_timeline(wb, ranges, demo):
    cols = [
        Col("Time", 10, "h:mm AM/PM", align="center"), Col("Event", 26, wrap=True),
        Col("Location", 20), Col("Person Responsible", 18), Col("Vendor", 20), Col("Notes", 24, wrap=True),
    ]
    demo_rows = []
    for time, event, loc, person, vendor, notes in D.TIMELINE:
        demo_rows.append({"Event": event, "Location": loc,
                           "Person Responsible": person, "Vendor": vendor, "Notes": notes})
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Timeline", "WEDDING DAY TIMELINE", "Minute-by-minute plan for the big day.",
        cols, demo_rows, "TimelineTable", 16, demo, nav_links=NAV_HOME)
    # convert time strings to proper time fractions
    if demo:
        for i, (time, *_rest) in enumerate(D.TIMELINE):
            r = first + i
            hh, mm = time.split(":")
            frac = (int(hh) * 60 + int(mm)) / (24 * 60)
            ws.cell(row=r, column=2, value=frac)
    return ws, first, last


# ------------------------------------------------------------- GIFT TRACKER
def build_gift_tracker(wb, ranges, demo):
    cols = [
        Col("Guest Name", 20), Col("Gift", 26, wrap=True), Col("Gift Type", 15, validation_list=ranges["GiftType"], align="center"),
        Col("Value", 11, FMT_CURRENCY0, align="right"),
        Col("Thank You Sent", 14, validation_list=ranges["YesNo"], align="center"),
        Col("Date Received", 14, FMT_DATE, align="center"), Col("Notes", 20, wrap=True),
    ]
    demo_rows = D.build_gift_rows(26) if demo else []
    ws, hdr, first, last, lastcol = build_table_sheet(
        wb, "Gift Tracker", "GIFT TRACKER", "Track gifts received and thank-you notes sent.",
        cols, demo_rows, "GiftTable", 28, demo, nav_links=NAV_HOME)
    status_conditional_formatting(ws, f"F{first}:F{last}", {"Yes": LIGHT_SAGE, "No": MUTED_GOLD})
    srow = last + 2
    ws.cell(row=srow, column=2, value="GIFT SUMMARY").font = F_SECTION
    srow += 1
    add_summary_strip(ws, srow, lastcol, [
        ("Gifts Received", f'COUNTA(B{first}:B{last})', FMT_INT),
        ("Total Value", f"SUM(E{first}:E{last})", FMT_CURRENCY0),
        ("Thank-Yous Sent", f'COUNTIF(F{first}:F{last},"Yes")', FMT_INT),
        ("Thank-Yous Pending", f'COUNTIFS(B{first}:B{last},"<>",F{first}:F{last},"<>Yes")', FMT_INT),
    ])
    return ws, first, last
